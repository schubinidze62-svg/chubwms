import os
import re
import json
import shutil
import itertools
from datetime import datetime

import pandas as pd
from flask import Flask, render_template, request, jsonify
from werkzeug.utils import secure_filename

import sys
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

try:
    from openpyxl import load_workbook
except Exception:
    load_workbook = None


app = Flask(__name__, template_folder="templates", static_folder="static")

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# Cloud deployment: use DATA_DIR env var for persistent storage (Render disk)
# Locally: use BASE_DIR
DATA_DIR = os.environ.get("DATA_DIR", BASE_DIR)
os.makedirs(DATA_DIR, exist_ok=True)

LOCATIONS_FILE = os.path.join(DATA_DIR, "warehouse_locations.json")
LOCATIONS_BAK = os.path.join(DATA_DIR, "warehouse_locations.json.bak")
STOCK_FILE = os.path.join(DATA_DIR, "warehouse_stock.json")
STOCK_BAK = os.path.join(DATA_DIR, "warehouse_stock.json.bak")
CATALOG_FILE = os.path.join(DATA_DIR, "catalog.json")
CATALOG_BAK = os.path.join(DATA_DIR, "catalog.json.bak")
SCORES_FILE = os.path.join(DATA_DIR, "category_scores.txt")
SCORES_BAK = os.path.join(DATA_DIR, "category_scores.txt.bak")
UPLOAD_FOLDER = os.path.join(DATA_DIR, "uploads")
ALLOWED_EXTENSIONS = {"xlsx", "xlsm", "xls", "csv"}

CATALOG_BARCODE_COL = 2  # C
CATALOG_CATEGORY_COL = 6  # G

os.makedirs(UPLOAD_FOLDER, exist_ok=True)
app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER
app.config["MAX_CONTENT_LENGTH"] = 120 * 1024 * 1024


# ---------------- helpers ----------------
def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


def backup_file(path, bak_path):
    try:
        if os.path.isfile(path):
            shutil.copy2(path, bak_path)
    except Exception:
        pass


def read_json(path, default):
    if not os.path.isfile(path):
        return default
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return default


def write_json(path, bak_path, data):
    backup_file(path, bak_path)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def save_uploaded_file(file_storage):
    original_name = file_storage.filename or "upload.xlsx"
    ext = original_name.rsplit(".", 1)[1].lower() if "." in original_name else "xlsx"
    safe_name = secure_filename(original_name)
    base = safe_name.rsplit(".", 1)[0] if "." in safe_name else safe_name
    if not base:
        base = "upload"
    filename = f"{base}_{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}.{ext}"
    path = os.path.join(app.config["UPLOAD_FOLDER"], filename)
    file_storage.save(path)
    return path


def clean_text(value):
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass
    text = str(value).replace("\u00a0", " ").strip()
    if text.lower() in {"nan", "none", "null"}:
        return ""
    return text


def normalize_text(value):
    return re.sub(r"\s+", " ", clean_text(value).casefold())


def normalize_barcode(value):
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass

    text = str(value).strip()
    if not text:
        return ""

    if re.search(r"e\+", text.lower()):
        try:
            text = str(int(float(text)))
        except Exception:
            pass

    if text.endswith(".0"):
        text = text[:-2]

    text = text.replace("\u00a0", "")
    text = re.sub(r"\s+", "", text)
    text = text.replace("-", "").replace("_", "")

    digits = re.sub(r"\D", "", text)
    if len(digits) >= 6:
        return digits
    return text


def barcode_variants(value):
    b = normalize_barcode(value)
    if not b:
        return []
    out = [b]
    stripped = b.lstrip("0")
    if stripped and stripped not in out:
        out.append(stripped)
    if b.isdigit() and len(b) < 13:
        padded = b.zfill(13)
        if padded not in out:
            out.append(padded)
    return out


def safe_int(value, default=0):
    try:
        if value is None or pd.isna(value):
            return default
    except Exception:
        pass
    try:
        return int(float(str(value).replace(",", ".").strip()))
    except Exception:
        return default


def safe_float(value, default=0.0):
    try:
        if value is None or pd.isna(value):
            return default
    except Exception:
        pass
    try:
        return float(str(value).replace(",", ".").strip())
    except Exception:
        return default


def format_num(value):
    try:
        v = float(value)
        if v.is_integer():
            return str(int(v))
        return str(round(v, 2))
    except Exception:
        return str(value)


def natural_key(text):
    return [int(x) if x.isdigit() else x.lower() for x in re.split(r"(\d+)", str(text))]


def col_letter(zero_based_index):
    n = zero_based_index + 1
    letters = ""
    while n:
        n, rem = divmod(n - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


# ---------------- locations ----------------
def parse_loc_code(code):
    text = clean_text(code)
    m = re.search(r"^([01]?UR)_(\d+)_(\d+)_(\d+)", text)
    if m:
        prefix = m.group(1)
        bank_num = int(m.group(2))
        bay = int(m.group(3))
        level = int(m.group(4))
        full_bank = f"{prefix}_{bank_num}" if prefix != "UR" else f"UR_{bank_num}"
        return full_bank, bank_num, bay, level, m.group(0)
    
    # Generic format: e.g. R1_5_2
    m2 = re.match(r"^([A-Za-z0-9]+)[_\-\s]+(\d+)[_\-\s]+(\d+)$", text)
    if m2:
        return m2.group(1).upper(), 1, int(m2.group(2)), int(m2.group(3)), text

    return None, None, None, None, text


def guess_bank(location_name):
    full_bank, _, _, _, _ = parse_loc_code(location_name)
    if full_bank:
        return full_bank
    text = clean_text(location_name)
    if not text:
        return "GENERAL"
    m = re.match(r"^([A-Za-zა-ჰ]+)", text)
    if m:
        return m.group(1).upper()
    m = re.match(r"^(\d+)", text)
    if m:
        return f"BANK_{m.group(1)}"
    return "BANK"


def default_limit(code):
    code = clean_text(code).upper()
    if code.startswith("1UR"):
        return 200
    if code.startswith("0UR"):
        return 100
    if code.startswith("UR"):
        return 20
    return 50


def location_status(qty, score, limit):
    qty = float(qty or 0)
    score = float(score or 0)
    limit = float(limit or 0)
    if qty <= 0:
        return "gray"
    if limit > 0 and score > limit:
        return "red"
    if score <= 0:
        return "yellow"
    if limit > 0 and score >= limit * 0.70:
        return "yellow"
    return "green"


def status_label(status):
    return {
        "green": "კარგი",
        "yellow": "დასაკვირვებელი",
        "red": "გადასაწყობი",
        "gray": "ცარიელი",
    }.get(status, status)


def load_locations():
    data = read_json(LOCATIONS_FILE, [])
    if not isinstance(data, list):
        return []
    result = []
    for i, item in enumerate(data):
        if not isinstance(item, dict):
            continue
        name = clean_text(
            item.get("name")
            or item.get("Std_Code")
            or item.get("std_code")
            or item.get("LocationCode")
            or item.get("location")
            or item.get("loc")
        )
        if not name:
            continue
        full_bank, _, parsed_bay, parsed_level, std_code = parse_loc_code(name)
        # Prioritize parsed full_bank over generic "BANK 0" strings
        bank = full_bank or clean_text(item.get("bank") or item.get("Full_Bank")) or guess_bank(name)
        bay = parsed_bay if parsed_bay is not None else safe_int(item.get("bay") or item.get("Bay") or item.get("col"), i % 20)
        level = parsed_level if parsed_level is not None else safe_int(item.get("level") or item.get("Level") or item.get("row"), i // 20)
        result.append({
            "id": clean_text(item.get("id")) or std_code or name,
            "name": std_code or name,
            "std_code": std_code or name,
            "bank": bank,
            "bay": max(0, bay),
            "level": max(0, level),
            "row": max(0, level),
            "col": max(0, bay),
        })
    return result


def save_locations_store(locations_data):
    normalized = []
    for i, item in enumerate(locations_data or []):
        if not isinstance(item, dict):
            continue
        name = clean_text(
            item.get("name")
            or item.get("std_code")
            or item.get("Std_Code")
            or item.get("LocationCode")
            or item.get("location")
            or item.get("loc")
        )
        if not name:
            continue
        full_bank, _, parsed_bay, parsed_level, std_code = parse_loc_code(name)
        bank = full_bank or clean_text(item.get("bank") or item.get("Full_Bank")) or guess_bank(name)
        bay = parsed_bay if parsed_bay is not None else safe_int(item.get("bay") or item.get("Bay") or item.get("col"), i % 20)
        level = parsed_level if parsed_level is not None else safe_int(item.get("level") or item.get("Level") or item.get("row"), i // 20)
        normalized.append({
            "id": clean_text(item.get("id")) or std_code or name,
            "name": std_code or name,
            "std_code": std_code or name,
            "bank": bank,
            "bay": max(0, bay),
            "level": max(0, level),
            "row": max(0, level),
            "col": max(0, bay),
        })
    write_json(LOCATIONS_FILE, LOCATIONS_BAK, normalized)
    return True


# ---------------- stock / catalog ----------------
def load_stock():
    data = read_json(STOCK_FILE, {})
    result = {}
    if isinstance(data, dict):
        for key, qty in data.items():
            parts = str(key).split("||", 1)
            if len(parts) != 2:
                continue
            loc, barcode = parts
            loc = clean_text(loc)
            barcode = normalize_barcode(barcode)
            if loc and barcode:
                result[(loc, barcode)] = result.get((loc, barcode), 0) + safe_float(qty, 0)
    elif isinstance(data, list):
        for item in data:
            if not isinstance(item, dict):
                continue
            loc = clean_text(item.get("location") or item.get("Std_Code") or item.get("std_code"))
            barcode = normalize_barcode(item.get("barcode") or item.get("Barcode"))
            qty = safe_float(item.get("quantity") or item.get("qty") or item.get("Available_Qty"), 0)
            if loc and barcode:
                result[(loc, barcode)] = result.get((loc, barcode), 0) + qty
    return result


def save_stock(stock_dict):
    serial = {}
    for (loc, barcode), qty in stock_dict.items():
        loc = clean_text(loc)
        barcode = normalize_barcode(barcode)
        if loc and barcode:
            serial[f"{loc}||{barcode}"] = safe_float(qty, 0)
    write_json(STOCK_FILE, STOCK_BAK, serial)
    return True


def load_catalog():
    data = read_json(CATALOG_FILE, {})
    if not isinstance(data, dict):
        return {}
    result = {}
    for barcode, category in data.items():
        bc = normalize_barcode(barcode)
        cat = clean_text(category)
        if bc and cat:
            result[bc] = cat
    return result


def save_catalog(catalog_dict):
    write_json(CATALOG_FILE, CATALOG_BAK, catalog_dict)
    return True


# ---------------- scores ----------------
DEFAULT_SCORES_RAW = """
სათვალე 1
საათი 3
სანდალი 1
სუნამო 1
ჩექმა 1
ბათინკი 1
სპორტული ფეხსაცმელი 1
დახურული ფეხსაცმელი 1
შუზი 1
პოლო 1
ქურთუკი 6.5
ქუსლიანი ფეხსაცმელი 1
მაისური 1
ჩუსტი 1
კაბა 2
ჩანთა 4
ქუდი 1
საფულე 1
პალტო 10
შემოქმედებითი & განმავითარებელი სათამაშოები 2
შესაფუთი პროდუქცია 3
შარვალი 3
სარეკლამო პროდუქცია 1
სპორტული ზედა 2
ქამარი 1
შორტი 1
კომბინიზონი 5
პიჯაკი 10
ნაქსოვი 4
ტანსაცმლის კომპლექტი 4
ბიუსტჰალტერი 1
შარვალი ჯინსის 3
ქვედაბოლო 3
ესპადრელი 1
ჟილეტი 5
ქოლგა 5
ბლუზა 3
პერანგი 2
ხელთათმანი 1
ტრუსი 1
წინდა 0.5
შარფი 1
სამკაული 0.5
ბოდე 2
წინდა კომპლექტი 2
ღამის თეთრეული 3
შარვალი სპორტული 3
საცურაო თეთრეული 2
ტრუსი კომპლექტი 2
ბოდე კომპლექტი 2
"""


def parse_score_line(line):
    line = clean_text(line)
    if not line:
        return None
    if ":" in line:
        left, right = line.split(":", 1)
    elif "\t" in line:
        parts = line.split("\t")
        left = " ".join(parts[:-1])
        right = parts[-1]
    else:
        parts = line.rsplit(maxsplit=1)
        if len(parts) != 2:
            return None
        left, right = parts
    category = clean_text(left)
    score_text = clean_text(right)
    if not category:
        return None
    if score_text.upper() in {"", "N/A", "NA", "NAN", "NONE"}:
        return category, None
    score = safe_float(score_text, None)
    if score is None:
        return None
    return category, score


def parse_scores_text(text):
    scores = {}
    for line in str(text or "").splitlines():
        parsed = parse_score_line(line)
        if parsed:
            cat, score = parsed
            scores[cat] = score
    return scores


def load_default_scores():
    return parse_scores_text(DEFAULT_SCORES_RAW)


def load_scores():
    if not os.path.isfile(SCORES_FILE):
        return load_default_scores()
    scores = {}
    try:
        with open(SCORES_FILE, "r", encoding="utf-8") as f:
            for line in f:
                parsed = parse_score_line(line)
                if parsed:
                    cat, score = parsed
                    scores[cat] = score
    except Exception:
        scores = {}
    return scores or load_default_scores()


def load_scores_text():
    if not os.path.isfile(SCORES_FILE):
        scores = load_default_scores()
        return "\n".join(
            f"{cat}\t{format_num(score) if score is not None else 'N/A'}"
            for cat, score in sorted(scores.items(), key=lambda x: natural_key(x[0]))
        )
    try:
        with open(SCORES_FILE, "r", encoding="utf-8") as f:
            return f.read()
    except Exception:
        return ""


def save_scores_text(scores_dict):
    backup_file(SCORES_FILE, SCORES_BAK)
    with open(SCORES_FILE, "w", encoding="utf-8") as f:
        for cat in sorted(scores_dict.keys(), key=natural_key):
            score = scores_dict[cat]
            f.write(f"{cat}: {format_num(score) if score is not None else 'N/A'}\n")
    return True


# ---------------- excel readers ----------------
def looks_like_location(value):
    text = clean_text(value)
    if not text or len(text) > 120:
        return False
    full_bank, _, _, _, _ = parse_loc_code(text)
    if full_bank:
        return True
    low = text.lower()
    bad_words = {
        "barcode", "quantity", "qty", "category", "score", "price",
        "შტრიხკოდი", "რაოდენობა", "კატეგორია", "ქულა"
    }
    if low in bad_words:
        return False
    has_digit = any(ch.isdigit() for ch in text)
    has_letter = any(ch.isalpha() for ch in text)
    has_sep = any(sep in text for sep in ["_", "-", "/", "\\", "."])
    return has_digit and (has_letter or has_sep)


def load_locations_excel(filepath):
    try:
        df = pd.read_excel(filepath, dtype=str)
    except Exception as e:
        raise ValueError(f"Locations Excel ვერ წაიკითხა: {e}")
    if df.empty:
        raise ValueError("Locations Excel ცარიელია.")

    loc_col = None
    cols = {normalize_text(c): c for c in df.columns}
    for low, original in cols.items():
        if "locationcode" in low or "location code" in low or low == "location" or low == "loc" or "ლოკაცია" in low:
            loc_col = original
            break

    if loc_col is None:
        best_col = None
        best_hits = -1
        for col in df.columns:
            hits = int(df[col].apply(looks_like_location).sum())
            if hits > best_hits:
                best_hits = hits
                best_col = col
        if best_hits <= 0:
            raise ValueError("Locations Excel-ში LocationCode/ლოკაცია ვერ მოიძებნა.")
        loc_col = best_col

    locations = []
    seen = set()
    for i, value in enumerate(df[loc_col].tolist()):
        text = clean_text(value)
        if not looks_like_location(text):
            continue
        full_bank, _, parsed_bay, parsed_level, std_code = parse_loc_code(text)
        code = std_code or text
        if code in seen:
            continue
        seen.add(code)
        locations.append({
            "id": code,
            "name": code,
            "std_code": code,
            "bank": full_bank or guess_bank(code),
            "bay": parsed_bay if parsed_bay is not None else (i % 20),
            "level": parsed_level if parsed_level is not None else (i // 20),
            "row": parsed_level if parsed_level is not None else (i // 20),
            "col": parsed_bay if parsed_bay is not None else (i % 20),
        })

    if not locations:
        raise ValueError("Locations Excel-ში ლოკაციები ვერ მოიძებნა.")
    return locations


def normalize_int_text_series(s):
    return s.astype(str).str.strip().str.replace(r"\.0$", "", regex=True).str.replace(r"\s+", "", regex=True)


def build_inventory_std_code_best(df_inv, loc_codes):
    sample = df_inv.head(3000).copy()
    loc_codes = set([str(x).strip() for x in loc_codes if str(x).strip()])

    best_score = -1
    best_desc = ""
    best_series_full = None
    best_series_sample = None

    for col in df_inv.columns:
        s_sample = sample[col].astype(str).str.strip()
        mask = s_sample.str.match(r"^[01]?UR_\d+_\d+_\d+", na=False)
        if mask.sum() <= 0:
            continue
        extracted = s_sample.str.extract(r"^([01]?UR_\d+_\d+_\d+)", expand=False)
        std_sample = extracted.fillna(s_sample)
        score = int(std_sample.isin(loc_codes).sum())
        if score > best_score:
            best_score = score
            best_desc = f"direct location column: {col}, matched sample={score}"
            s_full = df_inv[col].astype(str).str.strip()
            extracted_full = s_full.str.extract(r"^([01]?UR_\d+_\d+_\d+)", expand=False)
            best_series_full = extracted_full.fillna(s_full)
            best_series_sample = std_sample

    prefix_cols, numeric_cols = [], []
    for col in df_inv.columns:
        col_norm = normalize_text(col)
        s = sample[col].dropna().astype(str).str.strip().head(1000)
        if s.empty:
            continue
        prefix_hits = s.str.upper().str.match(r"^[01]?UR$", na=False).sum()
        prefix_name = any(w in col_norm for w in ["aisle", "zone", "row", "რიგი", "ზონა"])
        if prefix_hits >= 5 or prefix_name:
            prefix_cols.append(col)
        numeric_clean = s.str.replace(r"\.0$", "", regex=True)
        numeric_hits = numeric_clean.str.match(r"^\d+$", na=False).sum()
        bad_name = any(w in col_norm for w in ["qty", "quantity", "available", "ნაშთი", "რაოდენობა", "barcode", "ბარკოდი", "sku", "შტრიხ"])
        if numeric_hits >= 5 and not bad_name:
            numeric_cols.append(col)

    preferred_prefix = [c for c in prefix_cols if any(w in normalize_text(c) for w in ["aisle", "row", "რიგი"])]
    if preferred_prefix:
        prefix_cols = preferred_prefix + [c for c in prefix_cols if c not in preferred_prefix]

    numeric_cols = numeric_cols[:14]
    prefix_cols = prefix_cols[:5]

    candidates = []
    if len(df_inv.columns) > 13:
        candidates.append((df_inv.columns[13], df_inv.columns[12], df_inv.columns[11], df_inv.columns[10], "old-index aisle/bank/bay/level"))

    for p in prefix_cols:
        for bank_c, bay_c, level_c in itertools.permutations(numeric_cols, 3):
            candidates.append((p, bank_c, bay_c, level_c, "auto-combo"))

    seen = set()
    for prefix_c, bank_c, bay_c, level_c, source in candidates:
        key = (prefix_c, bank_c, bay_c, level_c)
        if key in seen:
            continue
        seen.add(key)

        try:
            prefix_s = sample[prefix_c].astype(str).str.strip().str.upper()
            bank_s = normalize_int_text_series(sample[bank_c])
            bay_s = normalize_int_text_series(sample[bay_c])
            level_s = normalize_int_text_series(sample[level_c])

            std_sample = prefix_s + "_" + bank_s + "_" + bay_s + "_" + level_s
            score = int(std_sample.isin(loc_codes).sum())

            if score > best_score:
                best_score = score
                best_desc = f"{source}: prefix={prefix_c}, bank={bank_c}, bay={bay_c}, level={level_c}, matched sample={score}"
                prefix_full = df_inv[prefix_c].astype(str).str.strip().str.upper()
                bank_full = normalize_int_text_series(df_inv[bank_c])
                bay_full = normalize_int_text_series(df_inv[bay_c])
                level_full = normalize_int_text_series(df_inv[level_c])
                best_series_full = prefix_full + "_" + bank_full + "_" + bay_full + "_" + level_full
                best_series_sample = std_sample
        except Exception:
            continue

    if best_series_full is None:
        raise ValueError("Inventory-ში ლოკაციის სვეტები ვერ ვიპოვე.")

    df_inv["Std_Code"] = best_series_full.astype(str).str.strip()
    valid_sample_matches = int(best_series_sample.isin(loc_codes).sum()) if best_series_sample is not None else 0

    if valid_sample_matches == 0:
        raise ValueError(
            "Inventory-ის ლოკაციები Locations ფაილს არ ემთხვევა.\n\n"
            f"არჩეული mapping:\n{best_desc}\n\n"
            "Inventory-ში Aisle/Bank/Bay/Level სვეტები სხვა რიგითაა ან სხვა სახელითაა."
        )

    return df_inv, best_desc, valid_sample_matches


def detect_qty_col(df):
    best_col = None
    best_score = -10**9

    exact_good_names = {
        "available qty", "available quantity", "available_qty",
        "qty available", "quantity", "qty", "on hand", "onhand",
        "stock", "available", "ნაშთი", "რაოდენობა", "ხელმისაწვდომი",
    }
    good_words = ["available", "qty", "quantity", "on hand", "onhand", "stock", "ნაშთი", "რაოდენობა", "ხელმისაწვდომი"]
    hard_bad_words = ["barcode", "bar code", "barcod", "ean", "upc", "gtin", "sku", "item", "product", "article", "id", "ბარკოდი", "შტრიხ", "კოდი"]
    location_bad_words = ["location", "loc", "aisle", "bank", "bay", "level", "ლოკაცია", "რიგი", "სექტორი", "სართული"]

    for col in df.columns:
        col_norm = normalize_text(col)
        sample = df[col].dropna().astype(str).head(1500)
        if sample.empty:
            continue

        normalized_for_barcode = sample.apply(normalize_barcode)
        barcode_like = int(normalized_for_barcode.apply(lambda x: bool(re.fullmatch(r"\d{8,18}", x))).sum())

        numeric = pd.to_numeric(
            sample.str.replace(",", ".", regex=False).str.replace(r"[^\d\.\-]", "", regex=True),
            errors="coerce",
        ).dropna()
        numeric_count = int(len(numeric))
        if numeric_count <= 0:
            continue

        numeric_sum = float(numeric.sum())
        numeric_mean = float(numeric.mean()) if numeric_count else 0
        numeric_max = float(numeric.max()) if numeric_count else 0

        has_good_name = any(w in col_norm for w in good_words)
        has_exact_good_name = col_norm in exact_good_names
        has_hard_bad = any(w in col_norm for w in hard_bad_words)
        has_location_bad = any(w in col_norm for w in location_bad_words)

        score = numeric_count
        if has_exact_good_name:
            score += 10000
        if has_good_name:
            score += 4000
        if numeric_sum > 0:
            score += 300
        if barcode_like >= max(20, numeric_count * 0.35) and not has_good_name:
            score -= 20000
        if has_hard_bad:
            score -= 25000
        if has_location_bad:
            score -= 8000
        if numeric_mean > 100000 and not has_good_name:
            score -= 15000
        if numeric_max > 10000000 and not has_good_name:
            score -= 12000

        if score > best_score:
            best_score = score
            best_col = col

    if best_col is None:
        raise ValueError("Inventory-ში Qty/Quantity/რაოდენობის სვეტი ვერ ვიპოვე.")
    return best_col


def detect_barcode_col(df):
    best_col = None
    best_score = -10**9
    good_words = ["barcode", "bar code", "ბარკოდი", "შტრიხ", "ean", "upc", "gtin"]
    bad_words = ["location", "loc", "aisle", "bank", "bay", "level", "qty", "quantity", "available", "ნაშთი", "რაოდენობა", "ლოკაცია", "რიგი", "სექტორი", "სართული"]

    for col in df.columns:
        col_norm = normalize_text(col)
        sample = df[col].dropna().astype(str).head(800)
        if sample.empty:
            continue

        normalized = sample.apply(normalize_barcode)
        barcode_like = int(normalized.apply(lambda x: bool(re.fullmatch(r"\d{8,18}", x))).sum())
        unique_count = int(normalized.nunique(dropna=True))
        score = barcode_like * 5 + min(unique_count, 800)

        if any(w in col_norm for w in good_words):
            score += 5000
        if any(w in col_norm for w in bad_words):
            score -= 5000

        if score > best_score:
            best_score = score
            best_col = col

    if not best_col:
        raise ValueError("Inventory-ში ბარკოდის სვეტი ვერ ვიპოვე.")
    return best_col


def load_inventory_excel(filepath, loc_codes):
    try:
        df_inv = pd.read_excel(filepath, dtype=str)
    except Exception as e:
        raise ValueError(f"Inventory Excel ვერ წაიკითხა: {e}")

    if df_inv.empty:
        return {}, set(), "Inventory ცარიელია"

    df_inv, inv_loc_debug, _ = build_inventory_std_code_best(df_inv, loc_codes)

    qty_col = detect_qty_col(df_inv)
    qty_raw = df_inv[qty_col].astype(str).str.replace(",", ".", regex=False).str.replace(r"[^\d\.\-]", "", regex=True)
    df_inv["Available_Qty"] = pd.to_numeric(qty_raw, errors="coerce").fillna(0)

    qty_total_check = float(df_inv["Available_Qty"].sum())
    qty_max_check = float(df_inv["Available_Qty"].max()) if len(df_inv) else 0
    if qty_total_check > 10000000 or qty_max_check > 1000000:
        raise ValueError(
            "Inventory-ში Qty სვეტი სავარაუდოდ არასწორად ამოიცნო.\n\n"
            f"არჩეული Qty column: {qty_col}\n"
            f"Qty sum: {format_num(qty_total_check)}\n"
            f"Qty max: {format_num(qty_max_check)}\n\n"
            "Inventory Excel-ში რაოდენობის სვეტის სათაურს გადაარქვით ზუსტად: Available Qty"
        )

    barcode_col = detect_barcode_col(df_inv)
    df_inv["Barcode"] = df_inv[barcode_col].apply(normalize_barcode)

    positive_qty_rows = int((df_inv["Available_Qty"] > 0).sum())
    matched_location_rows = int(df_inv["Std_Code"].astype(str).isin(loc_codes).sum())

    if positive_qty_rows > 0 and matched_location_rows == 0:
        raise ValueError(
            "Inventory-ში რაოდენობა არის, მაგრამ არცერთი ლოკაცია არ დაემთხვა Locations ფაილს.\n\n"
            f"Inventory mapping: {inv_loc_debug}\n"
            f"Qty column: {qty_col}\n"
            f"Barcode column: {barcode_col}"
        )

    stock = {}
    for _, row in df_inv.iterrows():
        loc = clean_text(row.get("Std_Code"))
        barcode = normalize_barcode(row.get("Barcode"))
        qty = safe_float(row.get("Available_Qty"), 0)
        if not loc or not barcode:
            continue
        stock[(loc, barcode)] = stock.get((loc, barcode), 0) + qty

    needed_keys = set()
    for (_, barcode), qty in stock.items():
        if qty > 0:
            for v in barcode_variants(barcode):
                needed_keys.add(v)

    debug = (
        f"Inventory: rows={len(df_inv)}, stock={len(stock)}, positive_qty_rows={positive_qty_rows}, "
        f"matched_location_rows={matched_location_rows}, mapping=({inv_loc_debug}), qty_col={qty_col}, barcode_col={barcode_col}"
    )
    return stock, needed_keys, debug


def load_catalog_csv_subset(filepath, needed_keys):
    bc_col = CATALOG_BARCODE_COL
    cat_col = CATALOG_CATEGORY_COL
    encodings = ["utf-8", "utf-8-sig"]
    last_error = None

    for enc in encodings:
        try:
            catalog_map = {}
            found = 0
            total = 0
            for chunk in pd.read_csv(filepath, dtype=str, header=None, usecols=[bc_col, cat_col], encoding=enc, chunksize=50000):
                chunk.columns = ["__BARCODE", "__CATEGORY"]
                for _, row in chunk.iterrows():
                    total += 1
                    barcode = normalize_barcode(row.get("__BARCODE", ""))
                    category = clean_text(row.get("__CATEGORY", ""))
                    if not barcode or not category:
                        continue
                    variants = barcode_variants(barcode)
                    if needed_keys and not any(v in needed_keys for v in variants):
                        continue
                    for v in variants:
                        catalog_map[v] = category
                    found += 1
                if needed_keys and needed_keys.issubset(set(catalog_map.keys())):
                    break
            return catalog_map, f"CSV C/G subset, scanned={total}, found={found}"
        except Exception as e:
            last_error = e

    raise ValueError(f"CSV კატალოგი ვერ წავიკითხე: {last_error}")


def load_catalog_excel_subset_fast(filepath, needed_keys):
    if load_workbook is None:
        raise ValueError("openpyxl არ არის დაყენებული. გაუშვით: python -m pip install openpyxl")

    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".xls":
        raise ValueError("Catalog ფაილი არის ძველი .xls ფორმატი. შეინახეთ Excel-ში .xlsx ან .xlsm-ად და თავიდან ატვირთეთ.")

    wb = load_workbook(filepath, read_only=True, data_only=True)
    min_col = min(CATALOG_BARCODE_COL, CATALOG_CATEGORY_COL) + 1
    max_col = max(CATALOG_BARCODE_COL, CATALOG_CATEGORY_COL) + 1
    bc_pos = CATALOG_BARCODE_COL + 1 - min_col
    cat_pos = CATALOG_CATEGORY_COL + 1 - min_col

    best_sheet = None
    best_score = -1
    sheet_logs = []

    try:
        for ws in wb.worksheets[:20]:
            barcode_hits = 0
            category_values = set()
            try:
                for row in ws.iter_rows(min_row=1, max_row=500, min_col=min_col, max_col=max_col, values_only=True):
                    barcode = normalize_barcode(row[bc_pos] if len(row) > bc_pos else "")
                    category = clean_text(row[cat_pos] if len(row) > cat_pos else "")
                    if barcode and re.fullmatch(r"\d{6,18}", barcode):
                        barcode_hits += 1
                    if category and not re.search(r"\d{6,18}", category):
                        category_values.add(category)
                score = barcode_hits * 10 + len(category_values)
                sheet_logs.append(f"{ws.title}: barcode={barcode_hits}, categories={len(category_values)}, score={score}")
                if score > best_score:
                    best_score = score
                    best_sheet = ws.title
                if barcode_hits >= 20 and len(category_values) >= 5:
                    best_sheet = ws.title
                    break
            except Exception as e:
                sheet_logs.append(f"{ws.title}: ERROR {e}")
    finally:
        wb.close()

    if not best_sheet:
        raise ValueError("კატალოგში C/G სვეტებით სწორი sheet ვერ ვიპოვე.")

    wb = load_workbook(filepath, read_only=True, data_only=True)
    ws = wb[best_sheet]
    catalog_map = {}
    found = 0
    scanned = 0

    try:
        for row in ws.iter_rows(min_row=1, min_col=min_col, max_col=max_col, values_only=True):
            scanned += 1
            barcode = normalize_barcode(row[bc_pos] if len(row) > bc_pos else "")
            category = clean_text(row[cat_pos] if len(row) > cat_pos else "")
            if not barcode or not category:
                continue
            if normalize_text(category) in {"category", "cat", "კატეგორია", "ჯგუფი"}:
                continue
            if not re.fullmatch(r"\d{6,18}", barcode):
                continue
            variants = barcode_variants(barcode)
            if needed_keys and not any(v in needed_keys for v in variants):
                continue
            for v in variants:
                catalog_map[v] = category
            found += 1
            if needed_keys and needed_keys.issubset(set(catalog_map.keys())):
                break
    finally:
        wb.close()

    debug = (
        f"Sheet={best_sheet}, C={col_letter(CATALOG_BARCODE_COL)}, G={col_letter(CATALOG_CATEGORY_COL)}, "
        f"scanned={scanned}, found={found}, map_keys={len(catalog_map)}"
    )

    if not catalog_map:
        raise ValueError(
            "კატალოგიდან Inventory-ის ბარკოდები ვერ მოიძებნა.\n\n"
            "შემოწმება: Catalog-ში ბარკოდი უნდა იყოს C სვეტში, კატეგორია G სვეტში.\n"
            f"Sheet check: {' | '.join(sheet_logs[:8])}"
        )

    return catalog_map, debug


def load_catalog_for_needed_barcodes(filepath, needed_keys):
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".csv":
        return load_catalog_csv_subset(filepath, needed_keys)
    return load_catalog_excel_subset_fast(filepath, needed_keys)


def map_category_for_barcode(catalog, barcode):
    for v in barcode_variants(barcode):
        if v in catalog:
            return catalog[v]
    return ""


# ---------------- report ----------------
def add_missing_inventory_locations(locations, stock):
    known = {loc["std_code"] for loc in locations}
    missing_codes = sorted({loc for (loc, bc) in stock.keys()} - known, key=natural_key)
    for code in missing_codes:
        full_bank, _, parsed_bay, parsed_level, std_code = parse_loc_code(code)
        if not std_code:
            continue
        locations.append({
            "id": std_code,
            "name": std_code,
            "std_code": std_code,
            "bank": full_bank or guess_bank(std_code),
            "bay": max(0, safe_int(parsed_bay, 0)),
            "level": max(0, safe_int(parsed_level, 0)),
            "row": max(0, safe_int(parsed_level, 0)),
            "col": max(0, safe_int(parsed_bay, 0)),
        })
    return locations


def build_report():
    locations = load_locations()
    stock = load_stock()
    catalog = load_catalog()
    scores = load_scores()

    locations = add_missing_inventory_locations(locations, stock)

    location_score = {}
    location_qty = {}
    location_items = {}
    location_categories = {}
    location_top = {}
    unknown_categories = set()
    category_summary_by_loc = {}

    for (loc, barcode), qty in stock.items():
        category = map_category_for_barcode(catalog, barcode)
        score_per_unit = scores.get(category, 0) if category else 0

        if barcode and not category:
            unknown_categories.add("(ბარკოდი კატალოგში არაა)")
        elif category and category not in scores:
            unknown_categories.add(category)

        calculated = float(qty or 0) * float(score_per_unit or 0)
        location_score[loc] = location_score.get(loc, 0) + calculated
        location_qty[loc] = location_qty.get(loc, 0) + float(qty or 0)
        location_items[loc] = location_items.get(loc, 0) + 1

        if category:
            location_categories.setdefault(loc, set()).add(category)
            d = category_summary_by_loc.setdefault(loc, {})
            d[category] = d.get(category, 0) + float(qty or 0)

    for loc, cats in category_summary_by_loc.items():
        location_top[loc] = [{"category": k, "qty": v} for k, v in sorted(cats.items(), key=lambda x: x[1], reverse=True)[:3]]

    enriched_locations = []
    banks_summary = {}

    for loc in locations:
        code = loc["std_code"]
        qty = location_qty.get(code, 0)
        score = location_score.get(code, 0)
        limit = default_limit(code)
        status = location_status(qty, score, limit)

        item = dict(loc)
        item["qty"] = qty
        item["score"] = score
        item["limit"] = limit
        item["status"] = status
        item["status_label"] = status_label(status)
        item["items"] = location_items.get(code, 0)
        item["categories"] = sorted(list(location_categories.get(code, set())), key=natural_key)
        item["top_categories"] = location_top.get(code, [])
        enriched_locations.append(item)

        # Bank aggregation
        bank_name = loc.get("bank") or guess_bank(code)
        if bank_name not in banks_summary:
            banks_summary[bank_name] = {
                "name": bank_name,
                "total_locations": 0,
                "occupied": 0,
                "empty": 0,
                "total_qty": 0,
                "total_score": 0,
                "max_bay": 0,
                "min_bay": 999999,
                "max_level": 0,
                "min_level": 999999,
                "status_counts": {"green": 0, "yellow": 0, "red": 0, "gray": 0}
            }
        bs = banks_summary[bank_name]
        bs["total_locations"] += 1
        if qty > 0:
            bs["occupied"] += 1
        else:
            bs["empty"] += 1
        bs["total_qty"] += qty
        bs["total_score"] += score
        bs["status_counts"][status] = bs["status_counts"].get(status, 0) + 1
        
        bay = int(loc.get("bay") or 0)
        level = int(loc.get("level") or 0)
        if bay > bs["max_bay"]:
            bs["max_bay"] = bay
        if bay < bs["min_bay"]:
            bs["min_bay"] = bay
        if level > bs["max_level"]:
            bs["max_level"] = level
        if level < bs["min_level"]:
            bs["min_level"] = level

    for bs in banks_summary.values():
        if bs["min_bay"] == 999999:
            bs["min_bay"] = 0
        if bs["min_level"] == 999999:
            bs["min_level"] = 0

    banks = sorted(banks_summary.keys(), key=natural_key)
    occupied = len([loc for loc in enriched_locations if float(loc.get("qty") or 0) > 0])
    total_qty = sum(float(v or 0) for v in location_qty.values())

    return {
        "locations": enriched_locations,
        "banks": banks,
        "banks_summary": banks_summary,
        "stock_rows": len(stock),
        "total_qty": total_qty,
        "occupied": occupied,
        "empty_locations": max(0, len(enriched_locations) - occupied),
        "catalog_size": len(catalog),
        "scores": scores,
        "scores_text": load_scores_text(),
        "unknown_categories": sorted(unknown_categories, key=natural_key),
        "updated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }


# ---------------- routes ----------------
@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/data")
def api_data():
    return jsonify(build_report())


@app.route("/api/location/<path:loc_code>")
def api_location_detail(loc_code):
    loc_code = clean_text(loc_code)
    stock = load_stock()
    catalog = load_catalog()
    scores = load_scores()

    items = []
    total_qty = 0.0
    total_score = 0.0

    for (loc, barcode), qty in stock.items():
        if loc.strip().lower() == loc_code.lower():
            cat = map_category_for_barcode(catalog, barcode)
            score_per_unit = scores.get(cat, 0) if cat else 0
            sub_score = round(float(qty) * float(score_per_unit or 0), 2)
            total_qty += float(qty)
            total_score += sub_score
            items.append({
                "barcode": barcode,
                "category": cat or "(კატეგორიის გარეშე)",
                "qty": float(qty),
                "score_per_unit": score_per_unit,
                "subtotal_score": sub_score,
            })

    items.sort(key=lambda x: x["qty"], reverse=True)
    limit = default_limit(loc_code)
    status = location_status(total_qty, total_score, limit)

    return jsonify({
        "location": loc_code,
        "limit": limit,
        "status": status,
        "status_label": status_label(status),
        "total_qty": total_qty,
        "total_score": total_score,
        "items": items,
    })


@app.route("/api/build", methods=["POST"])
def api_build():
    messages = []
    try:
        locations_file = request.files.get("locations_file")
        inventory_file = request.files.get("inventory_file")
        catalog_file = request.files.get("catalog_file")
        scores_text = request.form.get("scores_text", "")

        if locations_file and locations_file.filename:
            if not allowed_file(locations_file.filename):
                return jsonify({"error": "Locations ფაილი უნდა იყოს Excel"}), 400
            path = save_uploaded_file(locations_file)
            locations = load_locations_excel(path)
            save_locations_store(locations)
            messages.append(f"ლოკაციები: {len(locations)}")

        parsed_scores = parse_scores_text(scores_text)
        if parsed_scores:
            save_scores_text(parsed_scores)
            messages.append(f"ქულები: {len(parsed_scores)} კატეგორია")

        locations = load_locations()
        loc_codes = {loc["std_code"] for loc in locations}

        if inventory_file and inventory_file.filename:
            if not allowed_file(inventory_file.filename):
                return jsonify({"error": "Inventory ფაილი უნდა იყოს Excel"}), 400
            if not loc_codes:
                return jsonify({"error": "ჯერ ატვირთეთ Locations Excel ლოკაციები tab-ში."}), 400
            path = save_uploaded_file(inventory_file)
            stock, needed_keys, inv_debug = load_inventory_excel(path, loc_codes)
            save_stock(stock)
            messages.append(f"ინვენტარი: {len(stock)} პოზიცია")
            messages.append(inv_debug)
        else:
            stock = load_stock()
            needed_keys = set()
            for (_, barcode), qty in stock.items():
                if float(qty or 0) > 0:
                    for v in barcode_variants(barcode):
                        needed_keys.add(v)

        if catalog_file and catalog_file.filename:
            if not allowed_file(catalog_file.filename):
                return jsonify({"error": "Catalog ფაილი უნდა იყოს Excel/CSV"}), 400
            if not needed_keys:
                return jsonify({"error": "Catalog-ისთვის ჯერ Inventory ატვირთეთ ან existing stock უნდა არსებობდეს."}), 400
            path = save_uploaded_file(catalog_file)
            catalog, cat_debug = load_catalog_for_needed_barcodes(path, needed_keys)
            save_catalog(catalog)
            messages.append(f"კატალოგი: {len(catalog)} key")
            messages.append(cat_debug)

        report = build_report()
        report["messages"] = messages or ["მონაცემები განახლდა"]
        return jsonify(report)

    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/save_map", methods=["POST"])
def save_map():
    try:
        data = request.get_json()
        if data is None or "locations" not in data:
            return jsonify({"error": "არასწორი მონაცემები"}), 400
        save_locations_store(data["locations"])
        return jsonify({"status": "ok"})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/save_scores_manual", methods=["POST"])
def save_scores_manual():
    try:
        data = request.get_json()
        if not isinstance(data, dict):
            return jsonify({"error": "არასწორი ფორმატი"}), 400

        scores = {}
        for cat, value in data.items():
            cat = clean_text(cat)
            if not cat:
                continue
            if value is None or str(value).upper() in {"N/A", "NA", "NONE"}:
                scores[cat] = None
            else:
                score = safe_float(value, None)
                if score is not None:
                    scores[cat] = score

        save_scores_text(scores)
        return jsonify({"status": "ok"})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/api/sample-data", methods=["POST"])
def api_sample_data():
    """Generates realistic demo warehouse data for testing."""
    try:
        scores = load_default_scores()
        save_scores_text(scores)

        banks_config = [
            ("1UR", 1, 10, 4),
            ("1UR", 2, 10, 4),
            ("0UR", 1, 8, 3),
            ("UR", 1, 6, 2),
        ]
        sample_locs = []
        for prefix, b_num, max_bay, max_level in banks_config:
            bank_name = f"{prefix}_{b_num}" if prefix != "UR" else f"UR_{b_num}"
            for bay in range(1, max_bay + 1):
                for level in range(1, max_level + 1):
                    code = f"{prefix}_{b_num}_{bay}_{level}"
                    sample_locs.append({
                        "id": code,
                        "name": code,
                        "std_code": code,
                        "bank": bank_name,
                        "bay": bay,
                        "level": level,
                        "row": level,
                        "col": bay,
                    })
        save_locations_store(sample_locs)

        catalog = {
            "486000100001": "ქურთუკი",
            "486000100002": "პიჯაკი",
            "486000100003": "მაისური",
            "486000100004": "შარვალი",
            "486000100005": "შარვალი ჯინსის",
            "486000100006": "ჩანთა",
            "486000100007": "ფეხსაცმელი",
            "486000100008": "წინდა",
            "486000100009": "პალტო",
            "486000100010": "კაბა",
        }
        save_catalog(catalog)

        stock = {}
        stock[("1UR_1_1_1", "486000100001")] = 15
        stock[("1UR_1_1_2", "486000100002")] = 25
        stock[("1UR_1_2_1", "486000100003")] = 160
        stock[("1UR_1_2_2", "486000100004")] = 30
        stock[("1UR_1_3_1", "486000100005")] = 40
        stock[("1UR_1_4_2", "486000100009")] = 22
        stock[("1UR_2_2_1", "486000100006")] = 18
        stock[("1UR_2_5_3", "486000100010")] = 80
        stock[("0UR_1_2_1", "486000100003")] = 50
        stock[("0UR_1_3_2", "486000100004")] = 45
        stock[("UR_1_1_1", "486000100008")] = 20
        stock[("UR_1_2_1", "486000100006")] = 8
        save_stock(stock)

        report = build_report()
        report["messages"] = ["სადემონსტრაციო საწყობის მონაცემები წარმატებით ჩაიტვირთა!"]
        return jsonify(report)
    except Exception as e:
        return jsonify({"error": str(e)}), 500


def main():
    print("=====================================================")
    print("Warehouse WMS Map Server is Running!")
    print("Open in browser: http://127.0.0.1:5000")
    print("=====================================================")
    app.run(debug=False, host="0.0.0.0", port=5000, use_reloader=False)


if __name__ == "__main__":
    main()
